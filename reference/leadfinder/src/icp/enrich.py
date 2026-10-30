"""Build a full CRM-grade record for the top N leads.

Default N is 3, on purpose. PSI is effectively free (~25k/day) but each audit
is a slow server-side Lighthouse run, and Hunter verification is genuinely
scarce. Three is a day's worth of calls, not a list to admire.

Produces every field a CRM would hold: identity, contact, classification,
measured scores worst-to-best, ranked pain points with evidence, the selling
angle, estimated deal value, and a ready-to-send cold email and call opener.
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from .audit import PsiAudit
from .outreach import call_opener, cold_email
from .painpoints import band, pain_points, scores_worst_to_best
from .providers import HunterVerify, ProviderUnavailable, candidate_emails

FIELDS = [
    "priority", "tier", "business_name", "vertical", "business_type",
    "phone", "email_best", "email_status", "email_candidates",
    "address", "website", "has_website",
    "rating", "review_count",
    "score_performance", "score_seo", "score_accessibility", "score_best_practices",
    "lcp_seconds", "scores_worst_to_best",
    "need", "ability", "fit", "reach", "disqualified",
    "pain_1", "pain_1_evidence", "pain_1_angle", "pain_1_service", "pain_1_value",
    "pain_2", "pain_3", "all_pain_codes",
    "email_subject", "email_body", "call_opener",
    "audited_at", "source",
]

_TIER_LABEL = {100: "A", 72: "B", 48: "C", 10: "D"}

def enrich_one(lead: dict, ledger, verify_email: bool = False,
               do_audit: bool = True) -> dict:
    """Audit one lead and assemble the full record. Never raises on provider failure."""
    audit = None
    if do_audit and lead.get("has_website") and lead.get("website"):
        try:
            audit = PsiAudit(ledger).audit(lead["website"])
        except Exception as e:
            audit = None
            lead = dict(lead, _audit_error=f"{type(e).__name__}: {str(e)[:90]}")

    # Re-score with the audit we just ran. Previously the record shipped the
    # pre-audit guess (need=45) in the same row as the measured score
    # (performance=31); the two disagreed and nothing reconciled them.
    from .scoring import score_lead
    if audit:
        lead = {**lead, **score_lead(lead, audit)}

    pains = pain_points(lead, audit)
    worst = scores_worst_to_best(audit)
    cands = candidate_emails(lead.get("website") or "")

    email_best, email_status = (cands[0] if cands else ""), "unverified"
    if verify_email and cands:
        try:
            v = HunterVerify(ledger).verify(cands[0])
            email_status = v.get("status") or "unknown"
            if not v.get("sendable"):
                # first guess bounces: try the next pattern once, no further
                if len(cands) > 1:
                    v2 = HunterVerify(ledger).verify(cands[1])
                    if v2.get("sendable"):
                        email_best, email_status = cands[1], v2.get("status")
        except (ProviderUnavailable, Exception) as e:
            email_status = f"not verified ({type(e).__name__})"

    mail = cold_email(lead, audit, pains)
    rec = {
        "priority": lead.get("priority"), "tier": _TIER_LABEL.get(lead.get("fit"), "?"),
        "business_name": lead.get("name"), "vertical": lead.get("vertical"),
        "business_type": lead.get("place_type") or lead.get("vertical"),
        "phone": lead.get("phone"), "email_best": email_best, "email_status": email_status,
        "email_candidates": "; ".join(cands[:4]),
        "address": lead.get("address"), "website": lead.get("website") or "",
        "has_website": lead.get("has_website"),
        "rating": lead.get("rating"), "review_count": lead.get("review_count"),
        "score_performance": (audit or {}).get("performance"),
        "score_seo": (audit or {}).get("seo"),
        "score_accessibility": (audit or {}).get("accessibility"),
        "score_best_practices": (audit or {}).get("best_practices"),
        "lcp_seconds": round((audit or {}).get("lcp_ms", 0) / 1000, 1) if (audit or {}).get("lcp_ms") else None,
        "scores_worst_to_best": " < ".join(f"{k} {v}" for k, v in worst),
        "need": lead.get("need"), "ability": lead.get("ability"),
        "fit": lead.get("fit"), "reach": lead.get("reach"),
        "disqualified": lead.get("disqualified") or "",
        "pain_1": pains[0]["headline"] if pains else "none found",
        "pain_1_evidence": pains[0]["evidence"] if pains else "",
        "pain_1_angle": pains[0]["angle"] if pains else "",
        "pain_1_service": pains[0]["service"] if pains else "",
        "pain_1_value": pains[0]["est_value"] if pains else "",
        "pain_2": pains[1]["headline"] if len(pains) > 1 else "",
        "pain_3": pains[2]["headline"] if len(pains) > 2 else "",
        "all_pain_codes": ",".join(p["code"] for p in pains),
        "email_subject": mail["subject"], "email_body": mail["body"],
        "call_opener": call_opener(lead, audit, pains),
        "audited_at": datetime.now().isoformat(timespec="seconds"),
        "source": "google_places",
    }
    return rec

def export_csv(records: List[dict], path) -> Path:
    path = Path(path)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(records)
    return path

def export_xlsx(records: List[dict], path) -> Optional[Path]:
    """Formatted workbook: frozen header, colour-banded scores, wrapped email text."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    path = Path(path)
    wb = Workbook(); ws = wb.active; ws.title = "Leads"
    ws.append(FIELDS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    RED = PatternFill("solid", fgColor="F8CBAD")
    AMB = PatternFill("solid", fgColor="FFE699")
    GRN = PatternFill("solid", fgColor="C6E0B4")
    score_cols = [FIELDS.index(f) + 1 for f in
                  ("score_performance", "score_seo", "score_accessibility", "score_best_practices")]

    for rec in records:
        ws.append([rec.get(f) for f in FIELDS])
        r = ws.max_row
        for ci in score_cols:
            v = ws.cell(r, ci).value
            if isinstance(v, int):
                ws.cell(r, ci).fill = RED if v < 50 else AMB if v < 90 else GRN
        for f in ("email_body", "call_opener", "pain_1_evidence", "pain_1_angle"):
            ws.cell(r, FIELDS.index(f) + 1).alignment = Alignment(wrap_text=True, vertical="top")

    widths = {"business_name": 30, "address": 38, "website": 38, "email_body": 70,
              "call_opener": 60, "pain_1": 34, "pain_1_evidence": 60, "pain_1_angle": 50,
              "scores_worst_to_best": 34, "email_candidates": 34, "email_subject": 40}
    for i, f in enumerate(FIELDS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(f, 15)
    wb.save(path)
    return path
