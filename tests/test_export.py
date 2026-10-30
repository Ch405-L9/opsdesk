import datetime as dt

import pytest
from openpyxl import load_workbook

from opsdesk.export.xlsx import CHECK, ExportLocked, export, import_edits
from opsdesk.workbench import get_item, upsert_item, workbench
from opsdesk.workflow import set_fields


def _seed(conn):
    upsert_item(conn, "lead:a.example", {"name": "=HYPERLINK(\"http://evil\",\"x\")", "website": "https://a.example", "score": 70})
    upsert_item(conn, "lead:b.example", {"name": "Bravo", "phone": "+1 555 0100"})
    upsert_item(conn, "bid:usaspending:9", {"sub_agency": "Navy", "recipient": "Inc", "amount": 1000.0, "end": "2027-01-01", "reasons": ["a", "b"]})
    upsert_item(conn, "job:manual:1", {"title": "Dev", "company": "Co", "assessment": {"technology_alignment": 67, "gaps": ["Go"], "blockers": []}})


def _col(ws, title):
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value == title)


def test_workbook_structure_and_formatting(paths):
    out = paths.data_dir / "x.xlsx"
    with workbench(paths) as conn:
        _seed(conn)
        result = export(conn, out)
    assert result["counts"] == {"lead": 2, "bid": 1, "job": 1}
    wb = load_workbook(out)
    assert wb.sheetnames == ["Today", "Leads", "Bids", "Jobs", "_lists", "_meta"]
    assert wb["_lists"].sheet_state == "hidden"
    ws = wb["Leads"]
    assert ws.column_dimensions["A"].hidden and ws.freeze_panes == "C2"
    assert ws.cell(row=1, column=_col(ws, "Status")).fill.fgColor.rgb.endswith("FFF2CC")
    assert ws.cell(row=1, column=_col(ws, "Phone")).fill.fgColor.rgb.endswith("1F3864")
    lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("_lists!$A$2" in f for f in lists) and any("_lists!$E$2" in f for f in lists)
    assert sum(len(cf.rules) for cf in ws.conditional_formatting) >= 5
    names = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=2) for r in range(2, ws.max_row + 1)}
    evil = names['=HYPERLINK("http://evil","x")']
    assert evil.data_type == "s"  # stored as text, never a formula
    assert wb["Bids"].cell(row=2, column=_col(wb["Bids"], "End")).value == dt.datetime(2027, 1, 1)


def test_round_trip_applies_only_your_changes(paths):
    out = paths.data_dir / "rt.xlsx"
    with workbench(paths) as conn:
        _seed(conn)
        export(conn, out)
    wb = load_workbook(out)
    ws = wb["Leads"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "lead:b.example")
    ws.cell(row=row, column=_col(ws, "Status")).value = "contacted"
    ws.cell(row=row, column=_col(ws, "Done")).value = CHECK
    ws.cell(row=row, column=_col(ws, "Follow-up Date")).value = dt.datetime(2026, 11, 3)
    ws.cell(row=row, column=_col(ws, "Notes")).value = "called, left message"
    wb.save(out)
    with workbench(paths) as conn:
        set_fields(conn, "lead:a.example", status="meeting")  # newer edit made in the CLI
        rep = export(conn, out)
        a, b = get_item(conn, "lead:a.example"), get_item(conn, "lead:b.example")
    assert rep["imported_edits"] == 1
    assert (b["status"], b["done"], b["followup_date"]) == ("contacted", True, "2026-11-03")
    assert "called, left message" in b["notes"]
    assert a["status"] == "meeting"  # the stale workbook did not undo the CLI change
    assert (paths.data_dir / "rt.prev.xlsx").exists()


def test_notes_append_and_bad_values_reported(paths):
    out = paths.data_dir / "n.xlsx"
    with workbench(paths) as conn:
        _seed(conn)
        set_fields(conn, "job:manual:1", notes="first")
        export(conn, out)
    wb = load_workbook(out)
    ws = wb["Jobs"]
    ws.cell(row=2, column=_col(ws, "Notes")).value = "first\nsecond"
    ws.cell(row=2, column=_col(ws, "Status")).value = "not-a-status"
    wb.save(out)
    with workbench(paths) as conn:
        rep = import_edits(conn, out)
        job = get_item(conn, "job:manual:1")
    assert rep["errors"] and "not valid" in rep["errors"][0]
    assert job["status"] == "saved"


def test_refuses_when_workbook_is_open(paths):
    out = paths.data_dir / "lock.xlsx"
    with workbench(paths) as conn:
        _seed(conn)
        export(conn, out)
        (paths.data_dir / ".~lock.lock.xlsx#").write_text("x")
        with pytest.raises(ExportLocked):
            export(conn, out)
